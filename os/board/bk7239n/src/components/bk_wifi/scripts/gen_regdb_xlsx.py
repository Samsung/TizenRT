#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Parse wireless-regdb style db.txt -> channel-tiled Excel (2.4G+5G).

Usage:
    python3 gen_regdb_xlsx.py [db.txt] [output.xlsx]

With no arguments it reads ./db.txt (next to this script) and writes
./regdb_channel_map.xlsx.

Backend: uses openpyxl when available (best fidelity); otherwise falls back
to the bundled zero-dependency writer minixlsx.py, so a real styled .xlsx is
always produced with no third-party dependency. Force the bundled writer with
the environment variable REGDB_XLSX_BACKEND=mini.
"""
import os
import re
import sys
import math


# Backend selection: prefer openpyxl (best fidelity); if it is not installed,
# fall back to the bundled zero-dependency writer minixlsx so a real .xlsx is
# always produced. Set REGDB_XLSX_BACKEND=mini to force the bundled writer.
BACKEND = None
if os.environ.get("REGDB_XLSX_BACKEND", "").lower() != "mini":
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        BACKEND = "openpyxl"
    except ImportError:
        BACKEND = None
if BACKEND is None:
    from minixlsx import (Workbook, Font, PatternFill, Alignment, Border, Side,
                          get_column_letter)
    BACKEND = "minixlsx"

HERE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(HERE, "db.txt")
OUT = os.path.join(HERE, "regdb_channel_map.xlsx")

RULE_RE = re.compile(
    r'\(\s*([\d.]+)\s*-\s*([\d.]+)\s*@\s*([\d.]+)\s*\)\s*,\s*\(\s*([^)]*?)\s*\)(.*)'
)

GROUPS = {
    "2.4GHz":  "2.4~2.4835GHz",
    "UNII-1":  "UNII-1\n5.15-5.25GHz",
    "UNII-2A": "UNII-2A\n5.25-5.35GHz",
    "UNII-2C": "UNII-2C\n5.47-5.725GHz",
    "UNII-3":  "UNII-3\n5.725-5.85GHz",
    "UNII-4":  "UNII-4\n5.85-5.925GHz",
}


COUNTRY_NAMES = {
    "00": "World", "AD": "Andorra", "AE": "United Arab Emirates", "AF": "Afghanistan",
    "AI": "Anguilla", "AL": "Albania", "AM": "Armenia", "AN": "Netherlands Antilles",
    "AR": "Argentina", "AS": "American Samoa", "AT": "Austria", "AU": "Australia",
    "AW": "Aruba", "AZ": "Azerbaijan", "BA": "Bosnia and Herzegovina", "BB": "Barbados",
    "BD": "Bangladesh", "BE": "Belgium", "BF": "Burkina Faso", "BG": "Bulgaria",
    "BH": "Bahrain", "BL": "Saint Barthelemy", "BM": "Bermuda", "BN": "Brunei",
    "BO": "Bolivia", "BR": "Brazil", "BS": "Bahamas", "BT": "Bhutan", "BW": "Botswana",
    "BY": "Belarus", "BZ": "Belize", "CA": "Canada", "CF": "Central African Republic",
    "CH": "Switzerland", "CI": "Cote d'Ivoire", "CL": "Chile", "CN": "China",
    "CO": "Colombia", "CR": "Costa Rica", "CU": "Cuba", "CX": "Christmas Island",
    "CY": "Cyprus", "CZ": "Czechia", "DE": "Germany", "DK": "Denmark", "DM": "Dominica",
    "DO": "Dominican Republic", "DZ": "Algeria", "EC": "Ecuador", "EE": "Estonia",
    "EG": "Egypt", "ES": "Spain", "ET": "Ethiopia", "FI": "Finland", "FM": "Micronesia",
    "FR": "France", "GB": "United Kingdom", "GD": "Grenada", "GE": "Georgia",
    "GF": "French Guiana", "GH": "Ghana", "GL": "Greenland", "GP": "Guadeloupe",
    "GR": "Greece", "GT": "Guatemala", "GU": "Guam", "GY": "Guyana", "HK": "Hong Kong",
    "HN": "Honduras", "HR": "Croatia", "HT": "Haiti", "HU": "Hungary", "ID": "Indonesia",
    "IE": "Ireland", "IL": "Israel", "IN": "India", "IR": "Iran", "IS": "Iceland",
    "IT": "Italy", "JM": "Jamaica", "JO": "Jordan", "JP": "Japan", "KE": "Kenya",
    "KH": "Cambodia", "KN": "Saint Kitts and Nevis", "KP": "North Korea",
    "KR": "South Korea", "KW": "Kuwait", "KY": "Cayman Islands", "KZ": "Kazakhstan",
    "LB": "Lebanon", "LC": "Saint Lucia", "LI": "Liechtenstein", "LK": "Sri Lanka",
    "LS": "Lesotho", "LT": "Lithuania", "LU": "Luxembourg", "LV": "Latvia",
    "MA": "Morocco", "MC": "Monaco", "MD": "Moldova", "ME": "Montenegro",
    "MF": "Saint Martin", "MH": "Marshall Islands", "MK": "North Macedonia",
    "MN": "Mongolia", "MO": "Macao", "MP": "Northern Mariana Islands", "MQ": "Martinique",
    "MR": "Mauritania", "MT": "Malta", "MU": "Mauritius", "MV": "Maldives",
    "MW": "Malawi", "MX": "Mexico", "MY": "Malaysia", "NA": "Namibia", "NG": "Nigeria",
    "NI": "Nicaragua", "NL": "Netherlands", "NO": "Norway", "NP": "Nepal",
    "NZ": "New Zealand", "OM": "Oman", "PA": "Panama", "PE": "Peru",
    "PF": "French Polynesia", "PG": "Papua New Guinea", "PH": "Philippines",
    "PK": "Pakistan", "PL": "Poland", "PM": "Saint Pierre and Miquelon",
    "PR": "Puerto Rico", "PT": "Portugal", "PW": "Palau", "PY": "Paraguay",
    "QA": "Qatar", "RE": "Reunion", "RO": "Romania", "RS": "Serbia", "RU": "Russia",
    "RW": "Rwanda", "SA": "Saudi Arabia", "SE": "Sweden", "SG": "Singapore",
    "SI": "Slovenia", "SK": "Slovakia", "SN": "Senegal", "SR": "Suriname",
    "SV": "El Salvador", "SX": "Sint Maarten", "SY": "Syria",
    "TC": "Turks and Caicos Islands", "TD": "Chad", "TG": "Togo", "TH": "Thailand",
    "TN": "Tunisia", "TR": "Turkey", "TT": "Trinidad and Tobago", "TW": "Taiwan",
    "TZ": "Tanzania", "UA": "Ukraine", "UG": "Uganda", "US": "United States",
    "UY": "Uruguay", "UZ": "Uzbekistan", "VC": "Saint Vincent and the Grenadines",
    "VE": "Venezuela", "VI": "US Virgin Islands", "VN": "Vietnam", "VU": "Vanuatu",
    "WF": "Wallis and Futuna", "WS": "Samoa", "YE": "Yemen", "YT": "Mayotte",
    "ZA": "South Africa", "ZW": "Zimbabwe",
}


def parse_power(pstr):
    raw = pstr.strip()
    s = raw.replace(" ", "")
    m = re.match(r'([\d.]+)mW', s, re.IGNORECASE)
    if m:
        mw = float(m.group(1))
        dbm = 10.0 * math.log10(mw) if mw > 0 else 0.0
        return dbm, raw
    m = re.match(r'([\d.]+)', s)
    if m:
        return float(m.group(1)), raw
    return None, raw


def parse_db(path):
    countries = []
    cur = None
    with open(path, encoding="utf-8", errors="replace") as f:
        for line in f:
            stripped = line.strip()
            if stripped.startswith("#") or not stripped:
                continue
            if stripped.startswith("wmmrule"):
                cur = None
                continue
            if stripped.startswith("country"):
                m = re.match(r'country\s+(\S+?):\s*(\S*)', stripped)
                cur = {"code": m.group(1), "region": m.group(2) or "", "rules": []}
                countries.append(cur)
                continue
            if stripped.startswith("(") and cur is not None:
                body = stripped.split("#", 1)[0].strip()
                m = RULE_RE.match(body)
                if not m:
                    continue
                dbm, raw = parse_power(m.group(4))
                flags = []
                for tok in m.group(5).split(","):
                    tok = tok.strip()
                    if tok and not tok.startswith("wmmrule="):
                        flags.append(tok)
                cur["rules"].append({
                    "start": float(m.group(1)), "end": float(m.group(2)),
                    "dbm": dbm, "raw": raw, "flags": flags,
                })
    return countries


def c24(ch):
    return 2484 if ch == 14 else 2412 + (ch - 1) * 5


def build_channels():
    chans = []
    chans.append(("2.4GHz", "1~11", [c24(ch) for ch in range(1, 12)], 20))
    chans.append(("2.4GHz", "12~13", [c24(12), c24(13)], 20))
    chans.append(("2.4GHz", "14", [c24(14)], 20))
    for ch in [36, 40, 44, 48]:
        chans.append(("UNII-1", str(ch), [5000 + ch * 5], 20))
    for ch in [52, 56, 60, 64]:
        chans.append(("UNII-2A", str(ch), [5000 + ch * 5], 20))
    for ch in [100, 104, 108, 112, 116, 120, 124, 128, 132, 136, 140, 144]:
        chans.append(("UNII-2C", str(ch), [5000 + ch * 5], 20))
    for ch in [149, 153, 157, 161, 165]:
        chans.append(("UNII-3", str(ch), [5000 + ch * 5], 20))
    for ch in [169, 173, 177]:
        chans.append(("UNII-4", str(ch), [5000 + ch * 5], 20))
    return chans


def fmt_power(raw, dbm):
    s = raw.replace(" ", "")
    m = re.match(r'([\d.]+)mW', s, re.IGNORECASE)
    if m:
        return "%gmW" % float(m.group(1))
    return "%dmW" % round(10 ** (dbm / 10.0))


def rule_at(country, center, width):
    lo, hi = center - width / 2.0, center + width / 2.0
    best = None
    for r in country["rules"]:
        if lo >= r["start"] - 0.5 and hi <= r["end"] + 0.5:
            if best is None or (r["dbm"] or 0) < (best["dbm"] or 0):
                best = r
    return best


def eval_col(country, centers, width):
    best = None
    for c in centers:
        r = rule_at(country, c, width)
        if r is None:
            return None
        if best is None or (r["dbm"] or 0) < (best["dbm"] or 0):
            best = r
    return best


def main():
    db = sys.argv[1] if len(sys.argv) > 1 else DB
    out = sys.argv[2] if len(sys.argv) > 2 else OUT
    countries = parse_db(db)
    chans = build_channels()

    wb = Workbook()
    ws = wb.active
    ws.title = "信道功率平铺表"
    ws.sheet_view.showGridLines = False  # 关闭网格线，顶部留白更干净

    f_white = PatternFill("solid", fgColor="FFFFFF")
    f_yellow = PatternFill("solid", fgColor="FFFF00")
    f_red = PatternFill("solid", fgColor="FF0000")
    f_hdr = PatternFill("solid", fgColor="4472C4")
    f_chno = PatternFill("solid", fgColor="8EA9DB")
    f_meta = PatternFill("solid", fgColor="D6DCE4")
    FILLS = {"g": f_white, "y": f_yellow, "r": f_red, "x": f_white}

    thin = Side("thin", color="BBBBBB")
    med = Side("medium", color="404040")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_al = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_al = Alignment(horizontal="left", vertical="center", wrap_text=True)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_font2 = Font(bold=True, color="000000")
    data_font = Font(size=8)
    data_font_r = Font(size=8, bold=True, color="FFFFFF")

    # ---- layout offsets: leave top rows + a left spacer column ----
    COL0 = 2               # table starts at column B (A is a narrow spacer)
    R_GRP = 11             # table header (freq band) row; rows above = legend
    R_CHNO = R_GRP + 1
    DATA0 = R_GRP + 2

    meta_cols = ["No", "Code (Country)", "DFS\nRegion", "Bands"]
    n_meta = len(meta_cols)
    total_cols = n_meta + len(chans)
    first_col = COL0
    last_col = COL0 + total_cols - 1
    last_row = DATA0 + len(countries) - 1

    line_al = Alignment(horizontal="left", vertical="center", wrap_text=False)

    def wide(r, text, font, cspan=18):
        """One merged line spanning several columns (single line, no wrap)."""
        ws.merge_cells(start_row=r, start_column=COL0, end_row=r, end_column=COL0 + cspan)
        c = ws.cell(r, COL0, text)
        c.font = font
        c.alignment = line_al

    # ================= top title + legend =================
    wide(2, "Wi-Fi Regulatory Max EIRP by Channel (2.4GHz / 5GHz)", Font(bold=True, size=14))
    wide(3, "Source: os/board/bk7239n/src/components/bk_wifi/scripts/db.txt    |    Power: max EIRP in mW", Font(size=10), cspan=30)

    # legend box region
    LEG_R1, LEG_R2 = 5, 9
    LEG_C1, LEG_C2 = COL0, last_col
    SWATCH_TXT_END = COL0 + 8      # swatch text merge end (left block)
    NOTE_C1 = COL0 + 10            # notes start column (right block)

    # legend box: medium perimeter + thin horizontal lines between rows
    for rr in range(LEG_R1, LEG_R2 + 1):
        for cc in range(LEG_C1, LEG_C2 + 1):
            l = med if cc == LEG_C1 else None
            r = med if cc == LEG_C2 else None
            t = med if rr == LEG_R1 else thin
            b = med if rr == LEG_R2 else thin
            ws.cell(rr, cc).border = Border(left=l, right=r, top=t, bottom=b)

    # left block: title + color swatches
    ws.cell(LEG_R1, COL0, "Legend").font = Font(bold=True, size=12)
    swatches = [
        (6, f_white, "White: Allowed (Active Scan) / not covered (x)"),
        (7, f_yellow, "Yellow: DFS (radar channel; initial scan is Passive)"),
        (8, f_red, "Red: NO-IR (Passive scan only)"),
    ]
    for r, fill, text in swatches:
        sc = ws.cell(r, COL0 + 1, "")
        sc.fill = fill
        sc.border = box
        ws.merge_cells(start_row=r, start_column=COL0 + 2, end_row=r, end_column=SWATCH_TXT_END)
        t = ws.cell(r, COL0 + 2, text)
        t.font = Font(size=9)
        t.alignment = line_al

    # right block: notes laid out beside the swatches
    ws.cell(LEG_R1, NOTE_C1, "Notes:").font = Font(bold=True, size=10)
    notes = [
        "- Cell = channel max EIRP in mW;  'x' = not allowed / not covered in this country",
        "- Power unified to mW (dBm-source values converted via 10^(dBm/10))",
        "- 2.4G collapsed to 1~11 / 12~13 / 14; adjacent same-power channels merged",
        "- Active/Passive shown per-channel by color; 'DFS Region' = DFS domain (FCC/ETSI/JP)",
    ]
    for k, txt in enumerate(notes):
        rr = 6 + k
        ws.merge_cells(start_row=rr, start_column=NOTE_C1, end_row=rr, end_column=LEG_C2)
        cc = ws.cell(rr, NOTE_C1, txt)
        cc.font = Font(size=9)
        cc.alignment = line_al

    # ================= table header =================
    for i, name in enumerate(meta_cols):
        c = COL0 + i
        cell = ws.cell(R_GRP, c, name)
        ws.merge_cells(start_row=R_GRP, start_column=c, end_row=R_CHNO, end_column=c)
        cell.fill = f_hdr
        cell.font = hdr_font
        cell.alignment = center_al

    col = COL0 + n_meta
    for (grp, label, centers, width) in chans:
        cell = ws.cell(R_CHNO, col, label)
        cell.fill = f_chno
        cell.font = hdr_font2
        cell.alignment = center_al
        col += 1

    group_end_cols = set()
    start = COL0 + n_meta
    for idx in range(len(chans)):
        grp = chans[idx][0]
        nxt = chans[idx + 1][0] if idx + 1 < len(chans) else None
        if nxt != grp:
            end = COL0 + n_meta + idx
            g1 = ws.cell(R_GRP, start, GROUPS[grp])
            g1.fill = f_hdr
            g1.font = hdr_font
            g1.alignment = center_al
            if end > start:
                ws.merge_cells(start_row=R_GRP, start_column=start, end_row=R_GRP, end_column=end)
            group_end_cols.add(end)
            start = end + 1

    # ================= data =================
    rows_data = []
    for n, country in enumerate(countries, 1):
        row = DATA0 + n - 1
        code = country["code"]
        name = COUNTRY_NAMES.get(code, code)
        disp = "%s (%s)" % (code, name)
        region = country["region"] or "-"

        col_rules = [eval_col(country, ce, w) for (_, _, ce, w) in chans]
        bands = []
        if any(r for (g, _, _, _), r in zip(chans, col_rules) if g == "2.4GHz"):
            bands.append("2.4G")
        if any(r for (g, _, _, _), r in zip(chans, col_rules) if g != "2.4GHz"):
            bands.append("5G")

        for i, v in enumerate([n, disp, region, "/".join(bands)]):
            cell = ws.cell(row, COL0 + i, v)
            cell.alignment = left_al if i == 1 else center_al
            cell.fill = f_meta

        vals, codes = [], []
        col = COL0 + n_meta
        for (grp, label, centers, width), r in zip(chans, col_rules):
            cell = ws.cell(row, col)
            cell.alignment = center_al
            if r is None:
                val, cd = "x", "x"
            else:
                val = fmt_power(r["raw"], r["dbm"])
                if "NO-IR" in r["flags"]:
                    cd = "r"
                elif "DFS" in r["flags"]:
                    cd = "y"
                else:
                    cd = "g"
            cell.value = val
            cell.fill = FILLS[cd]
            cell.font = data_font_r if cd == "r" else data_font
            vals.append(val)
            codes.append(cd)
            col += 1
        rows_data.append((vals, codes))

    # borders over table region
    for rr in range(R_GRP, last_row + 1):
        for cc in range(first_col, last_col + 1):
            left = med if (cc == first_col or cc == COL0 + n_meta or (cc - 1) in group_end_cols) else thin
            right = med if (cc == last_col or cc == COL0 + n_meta - 1 or cc in group_end_cols) else thin
            top = med if (rr == R_GRP or rr == DATA0) else thin
            bottom = med if (rr == R_CHNO or rr == last_row) else thin
            ws.cell(rr, cc).border = Border(left=left, right=right, top=top, bottom=bottom)

    # merge consecutive identical channel cells within same group
    col0 = COL0 + n_meta
    for idx, (vals, codes) in enumerate(rows_data):
        row = DATA0 + idx
        i = 0
        while i < len(chans):
            j = i
            while (j + 1 < len(chans)
                   and chans[j + 1][0] == chans[i][0]
                   and vals[j + 1] == vals[i]
                   and codes[j + 1] == codes[i]):
                j += 1
            if j > i:
                ws.merge_cells(start_row=row, start_column=col0 + i,
                               end_row=row, end_column=col0 + j)
            i = j + 1

    # widths
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions[get_column_letter(COL0)].width = 4       # No
    ws.column_dimensions[get_column_letter(COL0 + 1)].width = 26  # Country (name + code)
    ws.column_dimensions[get_column_letter(COL0 + 2)].width = 9   # DFS
    ws.column_dimensions[get_column_letter(COL0 + 3)].width = 8   # Bands
    for c in range(COL0 + n_meta, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 6.5
    # widen the collapsed 2.4G "1~11" column (covers 11 channels)
    ws.column_dimensions[get_column_letter(COL0 + n_meta)].width = 11
    ws.row_dimensions[R_GRP].height = 30
    ws.row_dimensions[R_CHNO].height = 16

    wb.save(out)
    print("countries:", len(countries), "channel cols:", len(chans),
          "backend:", BACKEND, "-> saved:", out)


if __name__ == "__main__":
    main()
